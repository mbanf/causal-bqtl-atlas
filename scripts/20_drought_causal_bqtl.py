"""
20_drought_causal_bqtl.py — Extract drought ASE from Engelhorn supplement,
identify drought-causal bQTL, and compare with well-watered causal bQTL.

Data sources (all from MOESM5.xlsx):
  - Table_S10: Drought bQTL positions + FDR
  - Table_S13b: Drought haplotype-specific read counts (per hybrid, per gene, 3 reps)
  - Table_S6: Well-watered bQTL (already processed as bqtl_snp_ww.csv)
"""

import sys, warnings
import functools
print = functools.partial(print, flush=True)
warnings.filterwarnings('ignore')

import numpy as np
import pandas as pd
from pathlib import Path
from collections import Counter
from scipy import stats

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pipeline_config import config


def extract_drought_bqtl(moesm5_path):
    """Extract drought bQTL from Table_S10."""
    import openpyxl
    wb = openpyxl.load_workbook(moesm5_path, read_only=True)
    ws = wb['Table_S10']

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # SNP bQTL are in columns 0-2, INDEL in 4-6
    # Skip header rows (0=description, 1=SNP/INDEL, 2=Position/FDRp/Type)
    snp_bqtl = []
    for row in rows[3:]:
        pos_str = row[0]
        fdr = row[1]
        btype = row[2]
        if pos_str and fdr is not None:
            # Parse 'B73-chr1_457532' -> chr1, 457532
            parts = str(pos_str).replace('B73-', '')
            chrom, position = parts.split('_')
            snp_bqtl.append({
                'chr': chrom,
                'pos': int(position),
                'fdr_ds': float(fdr),
                'type': str(btype) if btype else 'Unknown',
            })

    df = pd.DataFrame(snp_bqtl)
    print(f"  Drought SNP bQTL extracted: {len(df)}")
    return df


def extract_drought_ase(moesm5_path):
    """Extract drought haplotype-specific read counts from Table_S13b."""
    import openpyxl
    wb = openpyxl.load_workbook(moesm5_path, read_only=True)
    ws = wb['Table_S13b']

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # Row 0: description
    # Row 1: hybrid names (B73xA188, B73xA619, ...) — every 6 columns
    # Row 2: column headers (Gene_B73, Rep1_B73, Rep2_B73, Rep3_B73, Rep1_NAM, Rep2_NAM, Rep3_NAM, ...)
    # Row 3+: data

    hybrid_row = rows[1]
    header_row = rows[2]

    # Parse hybrid names from row 1
    hybrids = []
    hybrid_cols = {}  # hybrid_name -> (start_col, end_col)
    current_hybrid = None
    start_col = None

    for i, val in enumerate(hybrid_row):
        if val and str(val).startswith('B73x'):
            if current_hybrid:
                hybrid_cols[current_hybrid] = (start_col, i)
            current_hybrid = str(val)
            start_col = i
            hybrids.append(current_hybrid)
    if current_hybrid:
        hybrid_cols[current_hybrid] = (start_col, len(hybrid_row))

    print(f"  Hybrids found: {len(hybrids)}")

    # For each hybrid, columns are: Rep1_B73, Rep2_B73, Rep3_B73, Rep1_NAM, Rep2_NAM, Rep3_NAM
    # Gene_B73 is in column 0

    results = []
    for row in rows[3:]:
        gene = row[0]
        if not gene or str(gene) == 'None':
            continue
        gene = str(gene)

        for hybrid in hybrids:
            start, end = hybrid_cols[hybrid]
            vals = row[start:end]

            # Extract B73 and NAM read counts (3 reps each)
            b73_reps = []
            nam_reps = []
            for j, v in enumerate(vals):
                if v is None or str(v) == 'NA' or str(v) == 'None':
                    continue
                try:
                    val = float(v)
                    if j < 3:
                        b73_reps.append(val)
                    else:
                        nam_reps.append(val)
                except (ValueError, TypeError):
                    continue

            if not b73_reps or not nam_reps:
                continue

            b73_mean = np.mean(b73_reps)
            nam_mean = np.mean(nam_reps)
            total = b73_mean + nam_mean

            if total < 10:  # minimum read depth
                continue

            # Avoid log(0)
            if b73_mean == 0 or nam_mean == 0:
                continue

            log2_ratio = np.log2(b73_mean / nam_mean)
            # Extract NAM parent name from hybrid
            nam_parent = hybrid.replace('B73x', '')

            results.append({
                'gene_id': gene,
                'hybrid': hybrid,
                'b73_mean': b73_mean,
                'nam_mean': nam_mean,
                'log2_ratio': log2_ratio,
                'abs_log2_ratio': abs(log2_ratio),
                'total_reads': total,
            })

    df = pd.DataFrame(results)
    print(f"  Drought ASE entries: {len(df)}")
    print(f"  Unique genes: {df['gene_id'].nunique()}")
    print(f"  Unique hybrids: {df['hybrid'].nunique()}")
    return df


def identify_causal_bqtl(bqtl_df, ase_df, gff3_path, promoter_size=2000, min_hybrids=3):
    """Identify causal bQTL using the same pipeline as well-watered.

    For each bQTL near a gene (within promoter_size of TSS):
      - Split hybrids into variant (has bQTL) vs reference (no bQTL)
      - Test if ASE is consistently different between groups
      - Compute Cohen's d as effect size
    """
    # Load genotype data to know which hybrids carry each bQTL
    geno = pd.read_csv(config.genotypes, sep='\t')
    hybrid_cols = [c for c in geno.columns if c not in ('chr', 'pos', 'ref')]

    # Load GFF3 for gene TSS positions
    print("  Loading GFF3 for TSS positions...")
    gff_genes = []
    with open(gff3_path) as f:
        for line in f:
            if line.startswith('#'):
                continue
            parts = line.strip().split('\t')
            if len(parts) < 9 or parts[2] != 'gene':
                continue
            chrom = parts[0]
            start = int(parts[3])
            end = int(parts[4])
            strand = parts[6]
            attrs = parts[8]
            gene_id = None
            for attr in attrs.split(';'):
                if attr.startswith('ID=gene:'):
                    gene_id = attr.split(':')[1]
                    break
                elif attr.startswith('ID=Zm'):
                    gene_id = attr[3:]
                    break
            if gene_id and chrom in config.valid_chroms:
                tss = start if strand == '+' else end
                gff_genes.append({'gene_id': gene_id, 'chr': chrom, 'tss': tss, 'strand': strand})

    genes_df = pd.DataFrame(gff_genes)
    print(f"  Genes with TSS: {len(genes_df)}")

    # For each bQTL, find nearby genes
    print("  Mapping bQTL to nearby genes...")
    # Filter bQTL to those in genotype data
    bqtl_set = set(zip(bqtl_df['chr'], bqtl_df['pos']))
    geno_set = set(zip(geno['chr'], geno['pos']))
    overlap = bqtl_set & geno_set
    print(f"  Drought bQTL in genotype data: {len(overlap)}/{len(bqtl_set)}")

    # Build gene-bQTL pairs
    pairs = []
    for _, gene in genes_df.iterrows():
        chrom = gene['chr']
        tss = gene['tss']
        # Find bQTL within promoter_size of TSS
        mask = (bqtl_df['chr'] == chrom) & \
               (bqtl_df['pos'] >= tss - promoter_size) & \
               (bqtl_df['pos'] <= tss + promoter_size)
        for _, bqtl in bqtl_df[mask].iterrows():
            if (bqtl['chr'], bqtl['pos']) in geno_set:
                pairs.append({
                    'chr': chrom,
                    'pos': bqtl['pos'],
                    'gene_id': gene['gene_id'],
                    'dist_to_tss': abs(bqtl['pos'] - tss),
                })

    pairs_df = pd.DataFrame(pairs)
    print(f"  bQTL-gene pairs within {promoter_size}bp: {len(pairs_df)}")

    # For each pair, compute Cohen's d
    print("  Computing Cohen's d for each pair...")
    causal_results = []
    ase_by_gene = ase_df.groupby('gene_id')

    for idx, pair in pairs_df.iterrows():
        chrom, pos, gene_id = pair['chr'], pair['pos'], pair['gene_id']

        # Get genotypes for this bQTL
        geno_row = geno[(geno['chr'] == chrom) & (geno['pos'] == pos)]
        if len(geno_row) == 0:
            continue
        geno_row = geno_row.iloc[0]
        ref_base = geno_row['ref']

        # Classify hybrids as variant or reference
        variant_hybrids = []
        ref_hybrids = []
        for h in hybrid_cols:
            val = str(geno_row[h])
            if val == '0|0' or val == ref_base:
                ref_hybrids.append(h)
            elif val != 'nan':
                variant_hybrids.append(h)

        if len(variant_hybrids) < min_hybrids or len(ref_hybrids) < min_hybrids:
            continue

        # Get ASE for this gene
        if gene_id not in ase_by_gene.groups:
            continue
        gene_ase = ase_by_gene.get_group(gene_id)

        # Get ASE values for variant and ref hybrids
        var_ase = gene_ase[gene_ase['hybrid'].isin(['B73x' + h for h in variant_hybrids])]['log2_ratio']
        ref_ase = gene_ase[gene_ase['hybrid'].isin(['B73x' + h for h in ref_hybrids])]['log2_ratio']

        if len(var_ase) < min_hybrids or len(ref_ase) < min_hybrids:
            continue

        # Cohen's d
        pooled_std = np.sqrt(((len(var_ase) - 1) * var_ase.std() ** 2 +
                              (len(ref_ase) - 1) * ref_ase.std() ** 2) /
                             (len(var_ase) + len(ref_ase) - 2))
        if pooled_std == 0:
            continue
        cohens_d = (var_ase.mean() - ref_ase.mean()) / pooled_std

        # Two-sample t-test
        t_stat, p_val = stats.ttest_ind(var_ase, ref_ase)

        causal_results.append({
            'chr': chrom,
            'pos': pos,
            'gene_id': gene_id,
            'cohens_d': cohens_d,
            'abs_d': abs(cohens_d),
            'p_value': p_val,
            'n_variant': len(var_ase),
            'n_ref': len(ref_ase),
            'variant_hybrids': '|'.join(variant_hybrids),
            'ref_hybrids': '|'.join(ref_hybrids),
        })

        if (idx + 1) % 5000 == 0:
            print(f"    Processed {idx+1}/{len(pairs_df)} pairs...")

    result_df = pd.DataFrame(causal_results)
    print(f"  Pairs tested: {len(result_df)}")

    # FDR correction
    from statsmodels.stats.multitest import multipletests
    if len(result_df) > 0:
        _, fdr, _, _ = multipletests(result_df['p_value'], method='fdr_bh')
        result_df['fdr'] = fdr
        causal = result_df[result_df['fdr'] < 0.05].copy()
        # Keep strongest per bQTL position
        causal = causal.sort_values('abs_d', ascending=False).drop_duplicates(['chr', 'pos'], keep='first')
        print(f"  Significant pairs (FDR<0.05): {len(result_df[result_df['fdr'] < 0.05])}")
        print(f"  Unique causal bQTL: {len(causal)}")
    else:
        causal = pd.DataFrame()

    return causal, result_df


def main():
    print("=" * 70)
    print("  DROUGHT CAUSAL bQTL PIPELINE")
    print("=" * 70)

    moesm5_path = config.moesm5

    # ── Step 1: Extract drought bQTL ──
    print("\n── Step 1: Extract drought bQTL from Table_S10 ──")
    ds_bqtl = extract_drought_bqtl(moesm5_path)
    ds_bqtl.to_csv(config.data_dir / 'processed' / 'bqtl_snp_ds.csv', index=False)
    print(f"  Saved: bqtl_snp_ds.csv")

    # ── Step 2: Extract drought ASE ──
    print("\n── Step 2: Extract drought ASE from Table_S13b ──")
    ds_ase = extract_drought_ase(moesm5_path)
    ds_ase.to_csv(config.data_dir / 'processed' / 'engelhorn_ase_ds.csv', index=False)
    print(f"  Saved: engelhorn_ase_ds.csv")

    # ── Step 3: Identify drought-causal bQTL ──
    print("\n── Step 3: Identify drought-causal bQTL ──")
    ds_causal, ds_all_pairs = identify_causal_bqtl(
        ds_bqtl, ds_ase, config.gff3, promoter_size=config.promoter_size
    )

    if len(ds_causal) == 0:
        print("  WARNING: No causal bQTL found under drought!")
        return

    ds_causal.to_csv(config.results_dir / 'causal_bqtl_ds.csv', index=False)
    print(f"  Saved: causal_bqtl_ds.csv ({len(ds_causal)} causal bQTL)")

    # ── Step 4: Compare WW vs DS ──
    print("\n" + "=" * 70)
    print("  COMPARISON: WELL-WATERED vs DROUGHT CAUSAL bQTL")
    print("=" * 70)

    ww_causal = pd.read_csv(config.results_dir / 'causal_bqtl_728.csv')
    print(f"\n  Well-watered causal bQTL: {len(ww_causal)}")
    print(f"  Drought causal bQTL: {len(ds_causal)}")

    # Overlap
    ww_keys = set(zip(ww_causal['chr'], ww_causal['pos']))
    ds_keys = set(zip(ds_causal['chr'], ds_causal['pos']))
    overlap = ww_keys & ds_keys
    ww_only = ww_keys - ds_keys
    ds_only = ds_keys - ww_keys

    print(f"\n  Overlap (causal in both): {len(overlap)} ({100*len(overlap)/len(ww_keys):.1f}% of WW)")
    print(f"  WW-only: {len(ww_only)}")
    print(f"  DS-only: {len(ds_only)}")

    # Compare effect sizes for overlapping bQTL
    if overlap:
        overlap_df = ww_causal[ww_causal.apply(lambda r: (r['chr'], r['pos']) in overlap, axis=1)].copy()
        ds_overlap = ds_causal[ds_causal.apply(lambda r: (r['chr'], r['pos']) in overlap, axis=1)].copy()

        # Merge on chr+pos
        merged = overlap_df[['chr', 'pos', 'gene_id', 'cohens_d', 'abs_d']].merge(
            ds_overlap[['chr', 'pos', 'cohens_d', 'abs_d']],
            on=['chr', 'pos'], suffixes=('_ww', '_ds')
        )

        rho, p = stats.spearmanr(merged['abs_d_ww'], merged['abs_d_ds'])
        print(f"\n  Effect size correlation (overlapping bQTL):")
        print(f"    Spearman rho = {rho:.3f}, p = {p:.2e}")
        print(f"    Median |d| WW = {merged['abs_d_ww'].median():.3f}")
        print(f"    Median |d| DS = {merged['abs_d_ds'].median():.3f}")

        # Direction consistency
        same_dir = (np.sign(merged['cohens_d_ww']) == np.sign(merged['cohens_d_ds'])).sum()
        print(f"    Same direction: {same_dir}/{len(merged)} ({100*same_dir/len(merged):.1f}%)")

    # Compare bQTL sets
    ww_bqtl_set = set(zip(pd.read_csv(config.bqtl)['chr'], pd.read_csv(config.bqtl)['pos']))
    ds_bqtl_set = set(zip(ds_bqtl['chr'], ds_bqtl['pos']))
    bqtl_overlap = ww_bqtl_set & ds_bqtl_set
    print(f"\n  Total bQTL overlap (WW∩DS): {len(bqtl_overlap)}/{len(ww_bqtl_set)} WW, {len(bqtl_overlap)}/{len(ds_bqtl_set)} DS")

    # Gene overlap
    ww_genes = set(ww_causal['gene_id'])
    ds_genes = set(ds_causal['gene_id'])
    gene_overlap = ww_genes & ds_genes
    print(f"  Gene overlap: {len(gene_overlap)}/{len(ww_genes)} WW, {len(gene_overlap)}/{len(ds_genes)} DS")

    # ── DS-only causal bQTL: drought-specific regulatory variants ──
    if ds_only:
        ds_only_df = ds_causal[ds_causal.apply(lambda r: (r['chr'], r['pos']) in ds_only, axis=1)]
        print(f"\n── Drought-specific causal bQTL ({len(ds_only_df)}) ──")
        print(f"  Median |d| = {ds_only_df['abs_d'].median():.3f}")
        print(f"  Unique genes: {ds_only_df['gene_id'].nunique()}")

        # Cross-reference with drought expression
        expr = pd.read_csv(config.expression, sep='\t')
        ds_only_expr = ds_only_df.merge(expr[['gene_id', 'log2fc_DS_vs_WW']], on='gene_id', how='left')
        drought_responsive = ds_only_expr[ds_only_expr['log2fc_DS_vs_WW'].abs() > 1]
        print(f"  Drought-responsive (|log2fc|>1): {len(drought_responsive)} bQTL, "
              f"{drought_responsive['gene_id'].nunique()} genes")

    # Save summary
    summary = {
        'n_ww_causal': len(ww_causal),
        'n_ds_causal': len(ds_causal),
        'n_overlap': len(overlap),
        'n_ww_only': len(ww_only),
        'n_ds_only': len(ds_only),
        'pct_ww_also_ds': 100 * len(overlap) / max(len(ww_keys), 1),
        'pct_ds_also_ww': 100 * len(overlap) / max(len(ds_keys), 1),
        'n_ds_bqtl_total': len(ds_bqtl),
        'n_ds_ase_genes': ds_ase['gene_id'].nunique(),
    }
    import json
    json.dump(summary, open(config.results_dir / 'ww_vs_ds_comparison.json', 'w'), indent=2)
    print(f"\n  Saved: ww_vs_ds_comparison.json")
    print("Done.")


if __name__ == '__main__':
    main()
